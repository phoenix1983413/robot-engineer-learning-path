#!/usr/bin/env python3
"""Convert Markdown to Word document with basic formatting."""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def add_markdown_to_docx(doc, text):
    """Add markdown text to docx with basic formatting."""
    lines = text.split('\n')
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        # Headings
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('## '):
            p = doc.add_heading(line[3:], level=2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('### '):
            p = doc.add_heading(line[4:], level=3)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('#### '):
            p = doc.add_heading(line[5:], level=4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r'^\d+\. ', line):
            content = re.sub(r'^\d+\. ', '', line)
            p = doc.add_paragraph(content, style='List Number')
        elif line.startswith('> '):
            p = doc.add_paragraph(line[2:])
            p.style = 'Intense Quote'
        elif line.startswith('---'):
            p = doc.add_paragraph('_' * 50)
        else:
            # Regular paragraph with inline formatting
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            # Parse inline markdown
            remaining = line
            while remaining:
                # Check for bold
                bold_match = re.search(r'\*\*(.+?)\*\*', remaining)
                # Check for italic
                italic_match = re.search(r'\*(.+?)\*', remaining)
                # Check for code
                code_match = re.search(r'`(.+?)`', remaining)
                
                # Find the earliest match
                matches = [m for m in [bold_match, italic_match, code_match] if m]
                if not matches:
                    # No more formatting, add remaining text
                    run = p.add_run(remaining)
                    run.font.size = Pt(12)
                    break
                    
                earliest = min(matches, key=lambda m: m.start())
                
                # Add text before match
                if earliest.start() > 0:
                    run = p.add_run(remaining[:earliest.start()])
                    run.font.size = Pt(12)
                
                # Add formatted text
                if earliest == bold_match:
                    run = p.add_run(earliest.group(1))
                    run.bold = True
                    run.font.size = Pt(12)
                elif earliest == italic_match:
                    run = p.add_run(earliest.group(1))
                    run.italic = True
                    run.font.size = Pt(12)
                elif earliest == code_match:
                    run = p.add_run(earliest.group(1))
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)
                
                remaining = remaining[earliest.end():]

def main():
    md_path = '/Users/deane/Documents/workspace-other/robot-engineer-learning-path/content.md'
    docx_path = '/Users/deane/Documents/workspace-other/robot-engineer-learning-path/机器人工程师学习路径.docx'
    
    with open(md_path, 'r') as f:
        text = f.read()
    
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    
    # Add title
    title = doc.add_heading('机器人工程师学习路径', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add subtitle
    subtitle = doc.add_paragraph('从入门到精通的完整知识体系')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()  # Empty line
    
    # Add content
    add_markdown_to_docx(doc, text)
    
    doc.save(docx_path)
    print(f'Saved to {docx_path}')

if __name__ == '__main__':
    main()
